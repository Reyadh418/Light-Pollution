"""
export_datasets.py
==================
Extraction and standardization pipeline for urban light pollution research datasets.
Processes the primary multi-sheet Excel workbook ("Database (Light Pollution).xlsx")
into FAIR-compliant open CSV files:
  1. data/dhaka_photometric_survey.csv (In-situ measurements across 13 urban sectors in Dhaka)
  2. data/benchtop_experiments.csv (Controlled benchtop luminaire & glazing experiments)

Authors: Md Reyad Hasan & Ishrat Jahan Riya
"""

import os
import re
import csv
import openpyxl

def parse_numeric(val):
    """Clean and parse numeric values, handling special edge cases and ranges."""
    if val is None:
        return None
    s = str(val).strip().replace(',', '')
    if s == '110/120':
        return 115.0  # Midpoint of 110-120 lx
    if s.endswith('r'):
        s = s[:-1]
    if s in ['?', '...', '', 'None']:
        return None
    try:
        return float(s)
    except ValueError:
        return None

def clean_str(val):
    """Normalize string, stripping whitespace and redundant newlines."""
    if val is None:
        return ''
    return ' '.join(str(val).split())

def export_photometric_survey(wb, output_path):
    """Extract and standardize 220 in-situ measurements across 13 urban sectors."""
    sector_mapping = {
        'Hospital': 'Hospital',
        'Roads': 'Roads',
        'Food and Restaurants': 'Food & Restaurants',
        'Fashion': 'Fashion',
        'School,Coaching and learning se': 'School, Coaching & Learning',
        'Buildings & Residence': 'Buildings & Residence',
        'Mobile and Accessories': 'Mobile & Accessories',
        'Variety shop': 'Variety Shop',
        'Cars and Transportation': 'Cars & Transportation',
        'Super shop and Convention hall': 'Super Shop & Convention Hall',
        'Visitable Places & Heritages': 'Visitable Places & Heritages',
        'Cinemas & entertainments': 'Cinemas & Entertainment',
        'Office': 'Office'
    }

    records = []
    record_counter = 1

    for sheetname in wb.sheetnames:
        if sheetname == 'Measurements from project':
            continue
        
        sheet = wb[sheetname]
        norm_sector = sector_mapping.get(sheetname, sheetname)
        
        # Check if first column is hospital header
        first_hdr = clean_str(sheet.cell(1, 1).value)
        is_hospital = 'Hospital Entrances Outdoors' in first_hdr
        
        last_known_area = ''
        
        for r in range(2, sheet.max_row + 1):
            if is_hospital:
                name_raw = sheet.cell(r, 2).value
                area_raw = sheet.cell(r, 3).value
                pos_raw = sheet.cell(r, 4).value
                color_raw = sheet.cell(r, 5).value
                low_raw = sheet.cell(r, 6).value
                high_raw = sheet.cell(r, 7).value
                avg_raw = sheet.cell(r, 8).value
            else:
                name_raw = sheet.cell(r, 1).value
                area_raw = sheet.cell(r, 2).value
                pos_raw = sheet.cell(r, 3).value
                color_raw = sheet.cell(r, 4).value
                low_raw = sheet.cell(r, 5).value
                high_raw = sheet.cell(r, 6).value
                avg_raw = sheet.cell(r, 7).value
            
            # Skip completely empty rows
            raw_vals = [name_raw, area_raw, pos_raw, color_raw, low_raw, high_raw, avg_raw]
            if not any(v is not None and str(v).strip() != '' for v in raw_vals):
                continue
                
            name = clean_str(name_raw)
            area = clean_str(area_raw)
            pos = clean_str(pos_raw)
            color = clean_str(color_raw)
            
            # Maintain hierarchical area if omitted in merged cells
            if area:
                last_known_area = area
            else:
                area = last_known_area
                
            low = parse_numeric(low_raw)
            high = parse_numeric(high_raw)
            avg = parse_numeric(avg_raw)
            
            # Skip category subheaders with no photometric data
            if low is None and high is None and avg is None:
                continue
                
            # Derive average if missing
            if avg is None:
                if low is not None and high is not None:
                    avg = round((low + high) / 2.0, 2)
                elif low is not None:
                    avg = low
                elif high is not None:
                    avg = high

            # Standardize color naming
            if color:
                color = color.replace('&', '+').replace('yellow', 'Yellow').replace('white', 'White')
            else:
                color = 'Unspecified'

            rec_id = f"DH-{norm_sector[:4].upper().replace(' ', '')}-{record_counter:03d}"
            record_counter += 1
            
            records.append({
                'record_id': rec_id,
                'sector': norm_sector,
                'site_name': name,
                'area': area,
                'measurement_position': pos or 'General Exterior',
                'light_color': color,
                'illuminance_lower_lx': f"{low:.1f}" if low is not None else '',
                'illuminance_higher_lx': f"{high:.1f}" if high is not None else '',
                'illuminance_average_lx': f"{avg:.1f}" if avg is not None else ''
            })

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        fieldnames = [
            'record_id', 'sector', 'site_name', 'area',
            'measurement_position', 'light_color',
            'illuminance_lower_lx', 'illuminance_higher_lx', 'illuminance_average_lx'
        ]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)

    print(f"Exported {len(records)} survey records to {output_path}")
    return records


def export_benchtop_experiments(wb, output_path):
    """Extract and standardize 12 controlled benchtop luminaire & glazing experiments."""
    sheet = wb['Measurements from project']
    
    experiment_metadata = [
        ('A1', 'Group A: Mounting Orientation & Vertical Elevation', 'Straight Vertical Reference', 5.0, 90, 'White', 2),
        ('A2', 'Group A: Mounting Orientation & Vertical Elevation', 'Angled Forward Tilt', 6.5, 45, 'White', 3),
        ('A3', 'Group A: Mounting Orientation & Vertical Elevation', 'Elevated Horizontal Cutoff (Taller)', 8.0, 0, 'White', 4),
        ('B1', 'Group B: LED Spectral Wavelength Impact', 'Cool White Emitter', 6.0, None, 'White', 5),
        ('B2', 'Group B: LED Spectral Wavelength Impact', 'Warm Yellow/Orange Emitter', 6.0, None, 'Yellow', 6),
        ('B3', 'Group B: LED Spectral Wavelength Impact', 'Short-Wavelength Blue Emitter', 6.0, None, 'Blue', 7),
        ('C1', 'Group C: Optical Cutoff Shielding Geometry', 'Unshielded Emitter (No cover)', 7.0, None, 'White', 8),
        ('C2', 'Group C: Optical Cutoff Shielding Geometry', 'Planar Shield (2D cover)', 7.0, None, 'White', 9),
        ('C3', 'Group C: Optical Cutoff Shielding Geometry', 'Conical/Angled Shield (Triangular cover)', 7.0, None, 'White', 10),
        ('C4', 'Group C: Optical Cutoff Shielding Geometry', 'Box Cutoff Shield (Square cover)', 7.0, None, 'White', 11),
        ('D1', 'Group D: Architectural Window Glazing Attenuation', 'Standard Transparent Window Glass', 2.0, None, 'White', 12),
        ('D2', 'Group D: Architectural Window Glazing Attenuation', 'Tinted/Attenuating Special Glass', 2.0, None, 'White', 13),
    ]

    records = []
    for eid, grp, desc, dist, tilt, default_color, row_idx in experiment_metadata:
        color_val = sheet.cell(row_idx, 4).value or default_color
        low_val = float(sheet.cell(row_idx, 5).value)
        avg_val = float(sheet.cell(row_idx, 6).value)
        high_val = float(sheet.cell(row_idx, 7).value)
        
        records.append({
            'experiment_id': eid,
            'experiment_group': grp,
            'configuration_tested': desc,
            'standoff_distance_inch': f"{dist:.1f}",
            'tilt_angle_deg': str(tilt) if tilt is not None else 'N/A',
            'light_color': color_val,
            'illuminance_min_lx': f"{low_val:.1f}",
            'illuminance_mean_lx': f"{avg_val:.1f}",
            'illuminance_peak_lx': f"{high_val:.1f}"
        })

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        fieldnames = [
            'experiment_id', 'experiment_group', 'configuration_tested',
            'standoff_distance_inch', 'tilt_angle_deg', 'light_color',
            'illuminance_min_lx', 'illuminance_mean_lx', 'illuminance_peak_lx'
        ]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)

    print(f"Exported {len(records)} benchtop experimental trials to {output_path}")
    return records


if __name__ == '__main__':
    # Determine base directory
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    
    # Try multiple potential paths for the raw workbook
    candidates = [
        os.path.join(base_dir, 'data', 'raw', 'Database (Light Pollution).xlsx'),
        os.path.join(base_dir, 'Documents', 'Database (Light Pollution).xlsx'),
        os.path.join(base_dir, 'Database (Light Pollution).xlsx')
    ]
    raw_excel = next((p for p in candidates if os.path.isfile(p)), None)
    if not raw_excel:
        raise FileNotFoundError(f"Could not find Database (Light Pollution).xlsx in candidate paths: {candidates}")
        
    print(f"Reading raw workbook from: {raw_excel}")
    wb = openpyxl.load_workbook(raw_excel, data_only=True)
    
    survey_csv = os.path.join(base_dir, 'data', 'dhaka_photometric_survey.csv')
    benchtop_csv = os.path.join(base_dir, 'data', 'benchtop_experiments.csv')
    
    export_photometric_survey(wb, survey_csv)
    export_benchtop_experiments(wb, benchtop_csv)
    print("Export complete.")

